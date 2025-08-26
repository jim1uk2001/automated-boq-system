from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, Protection
from openpyxl.worksheet.protection import SheetProtection
from typing import List, Dict, Any, Optional
import io
from ..models import BOQItem, Project
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from datetime import datetime

class ExcelGenerator:
    def __init__(self):
        self.header_font = Font(bold=True, size=12)
        self.section_font = Font(bold=True, size=14)
        self.title_font = Font(bold=True, size=18)
        self.data_font = Font(size=10)
        self.preamble_font = Font(size=10)
        
        self.vertical_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin')
        )
        
        self.full_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        self.header_fill = None
        self.section_fill = None
        self.rate_fill = None
        self.total_fill = None
        self.warning_fill = None

    def _get_measurement_standard_preambles(self, standard: str) -> Dict[str, Any]:
        """Get preambles for specific measurement standard"""
        
        smm7_preambles = {
            "general": [
                "All works shall be executed in accordance with the current edition of SMM7, relevant British Standards, Codes of Practice, and statutory requirements.",
                "All materials shall be new, of merchantable quality, and approved by the Architect/Contract Administrator."
            ],
            "measurement_rules": [
                "Quantities are measured net in accordance with SMM7 rules unless otherwise stated.",
                "Rates are deemed to include for supply and delivery of materials, labour, plant, tools, and equipment.",
                "Rates shall include for waste, laps, overlaps, cutting, jointing, notching, and normal fixings.",
                "No allowance shall be made for working in confined spaces, at heights, or in difficult conditions unless specifically described.",
                "All ancillary sundries, fittings, and supports necessary for proper execution are deemed included unless stated otherwise."
            ],
            "specific": {
                "plastering": "Rates include angles, arises, reveals, soffits, and openings under 0.5m².",
                "painting": "Rates include preparation, priming, knotting, stopping, rubbing down, cutting in, and scaffolding up to 3.6m height.",
                "brickwork_blockwork": "Rates include bonding, toothing, cutting, forming openings up to 0.5m², and all jointing unless otherwise described.",
                "concrete": "Rates include compacting, curing, finishing, formwork to sides up to 0.5m², and sundry supports unless otherwise measured.",
                "timber": "Rates include nails, screws, bolts, cutting, notching, drilling, and fitting of normal connectors."
            },
            "provisional_and_pc_sums": [
                "Prime Cost (PC) Sums are for supply of materials/goods only. Contractor's rates shall include for fixing, labour, and attendance.",
                "Provisional Sums are allowances for work not sufficiently defined; they are to be expended as directed by the Architect/Contract Administrator."
            ]
        }
        
        nrm2_preambles = {
            "general": [
                "All works shall comply with the latest British Standards, Codes of Practice, statutory requirements, and good building practice.",
                "All materials shall be new, of best quality of their kind, and approved by the Contract Administrator."
            ],
            "measurement_rules": [
                "Quantities measured net in accordance with RICS NRM2 unless otherwise stated.",
                "Rates shall include supply and delivery of materials, labour, plant, tools, equipment, waste, laps, overlaps, fixings, and ancillary items.",
                "Items measured in m² or m³ include all work in place unless otherwise described.",
                "Works at any level, height, or position are deemed included.",
                "All cutting, fitting, notching, splaying, jointing, and making good are deemed included."
            ],
            "specific": {
                "plastering": "Includes angles, arises, reveals, and openings <300mm wide.",
                "painting": "Includes cutting in, preparation, and all surfaces unless otherwise described.",
                "brickwork_blockwork": "Rates include bonding, quoins, reveals, jambs, sills, jointing, pointing, and cutting.",
                "concrete": "Rates include formwork, compacting, curing, finishing unless separately measured.",
                "timber": "Rates include nails, screws, plates, fixings."
            },
            "provisional_and_pc_sums": [
                "Prime Cost (PC) Sums cover supply-only of nominated items; contractor's rates include attendances, fixings, associated works.",
                "Provisional Sums must be identified as either Defined (scope broadly known) or Undefined (scope unknown)."
            ]
        }
        
        cesmm_preambles = {
            "general": [
                "All works shall be executed in accordance with CESMM4, relevant British Standards, and statutory requirements.",
                "All materials shall be new, of specified quality, and approved by the Engineer."
            ],
            "measurement_rules": [
                "Quantities measured net in accordance with CESMM4 unless otherwise stated.",
                "Rates deemed to include for all labour, materials, plant, and incidental costs.",
                "Method-related charges shall be included where applicable.",
                "All temporary works and construction aids are deemed included unless separately measured."
            ]
        }
        
        preambles_map = {
            "SMM7": smm7_preambles,
            "RICS_NRM": nrm2_preambles,
            "CESMM": cesmm_preambles
        }
        
        return preambles_map.get(standard, smm7_preambles)

    def generate_boq_excel(self, project: Project, boq_items: List[BOQItem]) -> bytes:
        """Generate Excel BOQ with protected cells, formulas, and drawing revision tracking"""
        wb = Workbook()
        
        drawing_ws = wb.active
        drawing_ws.title = "Drawing Register"
        
        drawing_ws['A1'] = f"PROJECT: {project.name}"
        drawing_ws['A1'].font = Font(bold=True, size=16)
        drawing_ws.merge_cells('A1:G1')
        
        drawing_ws['A2'] = f"MEASUREMENT STANDARD: {project.measurement_standard.value.upper()}"
        drawing_ws['A2'].font = Font(bold=True, size=12)
        drawing_ws.merge_cells('A2:G2')
        
        drawing_ws['A3'] = "DRAWING REGISTER - FOR VERSION CONTROL"
        drawing_ws['A3'].font = Font(bold=True, size=14)
        drawing_ws.merge_cells('A3:G3')
        
        drawing_headers = [
            "Drawing Number", "Drawing Title", "Revision", "Revision Date", 
            "Drawing Type", "File Name", "Upload Date"
        ]
        
        header_row = 5
        for col, header in enumerate(drawing_headers, 1):
            cell = drawing_ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True, size=12)
            cell.border = self.full_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        current_row = header_row + 1
        for drawing in project.drawings:
            drawing_ws.cell(row=current_row, column=1, value=drawing.drawing_number or "N/A")
            drawing_ws.cell(row=current_row, column=2, value=drawing.drawing_title or drawing.filename)
            drawing_ws.cell(row=current_row, column=3, value=drawing.revision_number or "A")
            drawing_ws.cell(row=current_row, column=4, value=drawing.revision_date.strftime("%Y-%m-%d") if drawing.revision_date else "N/A")
            drawing_ws.cell(row=current_row, column=5, value=drawing.drawing_type.value if drawing.drawing_type else "General")
            drawing_ws.cell(row=current_row, column=6, value=drawing.filename)
            drawing_ws.cell(row=current_row, column=7, value=drawing.uploaded_at.strftime("%Y-%m-%d"))
            
            for col in range(1, 8):
                drawing_ws.cell(row=current_row, column=col).border = self.vertical_border
            
            current_row += 1
        
        warning_row = current_row + 2
        drawing_ws.cell(row=warning_row, column=1, value="IMPORTANT: Contractors must verify they are pricing the correct drawing revisions listed above.")
        drawing_ws.cell(row=warning_row, column=1).font = Font(bold=True, size=12)
        drawing_ws.merge_cells(f'A{warning_row}:G{warning_row}')
        
        drawing_ws.cell(row=warning_row + 1, column=1, value="Any discrepancies between BOQ quantities and drawing revisions must be reported before bid submission.")
        drawing_ws.cell(row=warning_row + 1, column=1).font = Font(bold=True, size=12)
        drawing_ws.merge_cells(f'A{warning_row + 1}:G{warning_row + 1}')
        
        drawing_ws.cell(row=warning_row + 2, column=1, value="This register prevents disputes from superseded drawings that could lead to arbitration or litigation.")
        drawing_ws.cell(row=warning_row + 2, column=1).font = Font(bold=True, size=12)
        drawing_ws.merge_cells(f'A{warning_row + 2}:G{warning_row + 2}')
        
        column_widths = [15, 30, 10, 12, 15, 25, 12]
        for col, width in enumerate(column_widths, 1):
            column_letter = chr(64 + col)  # Convert column number to letter (A, B, C, etc.)
            drawing_ws.column_dimensions[column_letter].width = width
        
        for row in drawing_ws.iter_rows():
            for cell in row:
                cell.protection = Protection(locked=True)
        drawing_ws.protection.sheet = True
        drawing_ws.protection.password = "boq2024"
        
        ws = wb.create_sheet(title="Bill of Quantities")
        
        preambles = self._get_measurement_standard_preambles(project.measurement_standard.value)
        
        current_row = 1
        
        ws[f'A{current_row}'] = "BILL OF QUANTITIES"
        ws[f'A{current_row}'].font = self.title_font
        ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'A{current_row}:I{current_row}')
        current_row += 1
        
        ws[f'A{current_row}'] = f"Project: {project.name}"
        ws[f'A{current_row}'].font = Font(bold=True, size=14)
        current_row += 1
        
        project_info_from_drawings = self._extract_project_info_from_drawings(project.drawings)
        if project_info_from_drawings:
            if project_info_from_drawings.get('project_number'):
                ws[f'A{current_row}'] = f"Project Number: {project_info_from_drawings.get('project_number')}"
                ws[f'A{current_row}'].font = Font(bold=True, size=12)
                current_row += 1
            if project_info_from_drawings.get('client_name'):
                ws[f'A{current_row}'] = f"Client: {project_info_from_drawings.get('client_name')}"
                ws[f'A{current_row}'].font = Font(bold=True, size=12)
                current_row += 1
        
        ws[f'A{current_row}'] = f"Measurement Standard: {project.measurement_standard.value.upper()}"
        ws[f'A{current_row}'].font = Font(bold=True, size=12)
        current_row += 1
        
        ws[f'A{current_row}'] = f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws[f'A{current_row}'].font = Font(size=10)
        current_row += 2
        
        ws[f'A{current_row}'] = "PREAMBLES AND GENERAL CONDITIONS"
        ws[f'A{current_row}'].font = self.section_font
        ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'A{current_row}:I{current_row}')
        current_row += 1
        
        for preamble in preambles.get('general', []):
            ws[f'A{current_row}'] = f"• {preamble}"
            ws[f'A{current_row}'].font = self.preamble_font
            ws[f'A{current_row}'].alignment = Alignment(wrap_text=True)
            ws.merge_cells(f'A{current_row}:I{current_row}')
            current_row += 1
        
        current_row += 1
        
        ws[f'A{current_row}'] = "MEASUREMENT RULES"
        ws[f'A{current_row}'].font = self.section_font
        ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'A{current_row}:I{current_row}')
        current_row += 1
        
        for rule in preambles.get('measurement_rules', []):
            ws[f'A{current_row}'] = f"• {rule}"
            ws[f'A{current_row}'].font = self.preamble_font
            ws[f'A{current_row}'].alignment = Alignment(wrap_text=True)
            ws.merge_cells(f'A{current_row}:I{current_row}')
            current_row += 1
        
        current_row += 1
        
        if preambles.get('specific'):
            ws[f'A{current_row}'] = "SPECIFIC MEASUREMENT RULES"
            ws[f'A{current_row}'].font = self.section_font
            ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
            ws.merge_cells(f'A{current_row}:I{current_row}')
            current_row += 1
            
            for work_type, rule in preambles['specific'].items():
                ws[f'A{current_row}'] = f"{work_type.replace('_', ' ').title()}: {rule}"
                ws[f'A{current_row}'].font = self.preamble_font
                ws[f'A{current_row}'].alignment = Alignment(wrap_text=True)
                ws.merge_cells(f'A{current_row}:I{current_row}')
                current_row += 1
            
            current_row += 1
        
        if preambles.get('provisional_and_pc_sums'):
            ws[f'A{current_row}'] = "PROVISIONAL AND PRIME COST SUMS"
            ws[f'A{current_row}'].font = self.section_font
            ws[f'A{current_row}'].alignment = Alignment(horizontal='center', vertical='center')
            ws.merge_cells(f'A{current_row}:I{current_row}')
            current_row += 1
            
            for pc_rule in preambles['provisional_and_pc_sums']:
                ws[f'A{current_row}'] = f"• {pc_rule}"
                ws[f'A{current_row}'].font = self.preamble_font
                ws[f'A{current_row}'].alignment = Alignment(wrap_text=True)
                ws.merge_cells(f'A{current_row}:I{current_row}')
                current_row += 1
            
            current_row += 2
        
        headers = [
            "Item No.", "Item Code", "Description", "Unit", 
            "Quantity", "Rate", "Amount", "Category", "Trade"
        ]
        
        header_row = current_row
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = self.header_font
            cell.border = self.full_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.protection = Protection(locked=True)
        
        current_row = header_row + 1
        total_amount_formula_cells = []
        
        items_by_category = {}
        for item in boq_items:
            if item.category not in items_by_category:
                items_by_category[item.category] = []
            items_by_category[item.category].append(item)
        
        item_idx = 1
        section_number = 1
        
        for category, category_items in items_by_category.items():
            section_header_cell = ws.cell(row=current_row, column=1)
            section_header_cell.value = f"SECTION {section_number}: {category.upper()}"
            section_header_cell.font = Font(bold=True, size=12)
            section_header_cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.merge_cells(f'A{current_row}:I{current_row}')
            current_row += 1
            
            for item in category_items:
                if item.item_code == "SECTION":
                    cell = ws.cell(row=current_row, column=1)
                    cell.value = item.description
                    cell.font = Font(bold=True, size=12)
                    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
                    for col in range(1, 10):
                        cell = ws.cell(row=current_row, column=col)
                        cell.alignment = Alignment(horizontal='center')
                    current_row += 1
                    continue
                
                cell = ws.cell(row=current_row, column=1, value=item_idx)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.alignment = Alignment(horizontal='center')
                cell.protection = Protection(locked=True)
                
                cell = ws.cell(row=current_row, column=2, value=item.item_code)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.protection = Protection(locked=True)
                
                # Description
                cell = ws.cell(row=current_row, column=3, value=item.description)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.alignment = Alignment(wrap_text=True)
                cell.protection = Protection(locked=True)
                
                cell = ws.cell(row=current_row, column=4, value=item.unit)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.alignment = Alignment(horizontal='center')
                cell.protection = Protection(locked=True)
                
                cell = ws.cell(row=current_row, column=5, value=item.quantity)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right')
                cell.protection = Protection(locked=True)
                
                cell = ws.cell(row=current_row, column=6, value=0.00)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right')
                cell.protection = Protection(locked=False)  # UNLOCKED for rate entry
                
                amount_cell = ws.cell(row=current_row, column=7)
                amount_cell.value = f"=E{current_row}*F{current_row}"
                amount_cell.font = self.data_font
                amount_cell.border = self.vertical_border
                amount_cell.number_format = '0.00'
                amount_cell.alignment = Alignment(horizontal='right')
                amount_cell.protection = Protection(locked=True)
                total_amount_formula_cells.append(f"G{current_row}")
                
                cell = ws.cell(row=current_row, column=8, value=item.category)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.protection = Protection(locked=True)
                
                cell = ws.cell(row=current_row, column=9, value=item.trade)
                cell.font = self.data_font
                cell.border = self.vertical_border
                cell.protection = Protection(locked=True)
                
                current_row += 1
                item_idx += 1
            
            section_number += 1
        
        current_row += 1
        total_row = current_row
        
        total_label_cell = ws.cell(row=total_row, column=6, value="TOTAL:")
        total_label_cell.font = Font(bold=True, size=12)
        total_label_cell.alignment = Alignment(horizontal='right')
        total_label_cell.border = self.full_border
        total_label_cell.protection = Protection(locked=True)
        
        total_cell = ws.cell(row=total_row, column=7)
        if total_amount_formula_cells:
            total_cell.value = f"=SUM({','.join(total_amount_formula_cells)})"
        else:
            total_cell.value = 0.00
        total_cell.font = Font(bold=True, size=12)
        total_cell.number_format = '0.00'
        total_cell.alignment = Alignment(horizontal='right')
        total_cell.border = self.full_border
        total_cell.protection = Protection(locked=True)
        
        column_widths = [8, 12, 50, 8, 12, 12, 15, 15, 15]  # Wider description column
        for col, width in enumerate(column_widths, 1):
            column_letter = chr(64 + col)  # Convert column number to letter (A, B, C, etc.)
            ws.column_dimensions[column_letter].width = width
        
        for row in range(1, current_row + 1):
            cell = ws.cell(row=row, column=1)
            if cell.fill and cell.fill.start_color.rgb in ["FF366092", "FF4472C4"]:  # Header/section rows
                ws.row_dimensions[row].height = 25
            else:
                ws.row_dimensions[row].height = 18
        
        instruction_row = total_row + 3
        ws.cell(row=instruction_row, column=1, value="INSTRUCTIONS FOR CONTRACTORS:")
        ws.cell(row=instruction_row, column=1).font = Font(bold=True)
        
        instructions = [
            "1. Only the RATE column can be edited",
            "2. Enter your rates in the appropriate currency",
            "3. The AMOUNT column will calculate automatically",
            "4. Do not modify any other cells - they are protected",
            "5. Save the file and upload through the bidding portal"
        ]
        
        for i, instruction in enumerate(instructions):
            ws.cell(row=instruction_row + i + 1, column=1, value=instruction)
            ws.cell(row=instruction_row + i + 1, column=1).font = Font(size=10)
        
        ws.protection.sheet = True
        ws.protection.password = "boq2024"  # Set a password for protection
        ws.protection.formatCells = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.insertColumns = False
        ws.protection.insertRows = False
        ws.protection.deleteColumns = False
        ws.protection.deleteRows = False
        
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        return excel_buffer.getvalue()

    def _extract_project_info_from_drawings(self, drawings: List) -> Dict[str, str]:
        """Extract project information from processed drawings for BOQ front page"""
        project_info = {}
        
        for drawing in drawings:
            if hasattr(drawing, 'project_name') and drawing.project_name:
                project_info['project_name'] = drawing.project_name
            if hasattr(drawing, 'project_number') and drawing.project_number:
                project_info['project_number'] = drawing.project_number
            
            if hasattr(drawing, 'project_name') and drawing.project_name:
                import re
                client_match = re.search(r'(?:for\s+|client\s*:\s*)([A-Za-z\s&]+)', drawing.project_name, re.IGNORECASE)
                if client_match:
                    project_info['client_name'] = client_match.group(1).strip()
        
        return project_info
    
    def generate_boq_pdf(self, boq_items: List[Dict[str, Any]], project_name: str, 
                        drawings: List[Dict[str, Any]] = None, 
                        measurement_standard: str = "SMM7") -> bytes:
        """Generate a PDF version of the BOQ"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72,
                              topMargin=72, bottomMargin=18)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=12
        )
        
        story = []
        
        story.append(Paragraph(f"Bill of Quantities - {project_name}", title_style))
        story.append(Spacer(1, 12))
        
        story.append(Paragraph("Project Information", heading_style))
        project_info = [
            ['Project Name:', project_name],
            ['Measurement Standard:', measurement_standard],
            ['Generated:', datetime.now().strftime('%d/%m/%Y %H:%M')],
            ['Total Items:', str(len(boq_items))]
        ]
        
        project_table = Table(project_info, colWidths=[2*inch, 4*inch])
        project_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(project_table)
        story.append(Spacer(1, 20))
        
        if drawings:
            story.append(Paragraph("Drawing Register", heading_style))
            drawing_data = [['Drawing Title', 'Revision', 'Type', 'Status']]
            for drawing in drawings:
                drawing_data.append([
                    drawing.get('filename', 'N/A'),
                    drawing.get('revision', 'A'),
                    drawing.get('file_type', 'PDF').upper(),
                    drawing.get('status', 'Processed')
                ])
            
            drawing_table = Table(drawing_data, colWidths=[2.5*inch, 1*inch, 1*inch, 1.5*inch])
            drawing_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(drawing_table)
            story.append(PageBreak())
        
        story.append(Paragraph("Bill of Quantities", heading_style))
        
        trades = {}
        for item in boq_items:
            trade = item.get('trade', 'General')
            if trade not in trades:
                trades[trade] = []
            trades[trade].append(item)
        
        for trade, items in trades.items():
            story.append(Paragraph(f"{trade} Works", heading_style))
            
            boq_data = [['Item', 'Description', 'Unit', 'Quantity', 'Rate (£)', 'Amount (£)']]
            
            trade_total = 0
            for i, item in enumerate(items, 1):
                quantity = float(item.get('quantity', 0))
                rate = float(item.get('unit_rate', 0))
                amount = quantity * rate
                trade_total += amount
                
                boq_data.append([
                    f"{trade[0]}.{i:02d}",
                    item.get('description', 'N/A'),
                    item.get('unit', 'Nr'),
                    f"{quantity:.2f}",
                    f"{rate:.2f}",
                    f"{amount:.2f}"
                ])
            
            boq_data.append(['', f'Sub-total for {trade}', '', '', '', f"{trade_total:.2f}"])
            
            boq_table = Table(boq_data, colWidths=[0.8*inch, 2.5*inch, 0.8*inch, 1*inch, 1*inch, 1*inch])
            boq_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('ALIGN', (1, 1), (1, -2), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('BACKGROUND', (0, -1), (-1, -1), colors.lightgrey),
            ]))
            story.append(boq_table)
            story.append(Spacer(1, 20))
        
        grand_total = sum(float(item.get('quantity', 0)) * float(item.get('unit_rate', 0)) for item in boq_items)
        total_data = [['', '', '', '', 'GRAND TOTAL:', f"£{grand_total:.2f}"]]
        total_table = Table(total_data, colWidths=[0.8*inch, 2.5*inch, 0.8*inch, 1*inch, 1*inch, 1*inch])
        total_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 12),
            ('BACKGROUND', (0, 0), (-1, -1), colors.lightblue),
            ('GRID', (4, 0), (-1, -1), 2, colors.black),
        ]))
        story.append(total_table)
        
        story.append(Spacer(1, 30))
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            alignment=1
        )
        story.append(Paragraph("Generated by Bojim BOQ Production Software", footer_style))
        story.append(Paragraph(f"Compliant with {measurement_standard} Standards", footer_style))
        
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    def generate_bid_comparison_excel(self, project: Project, bids_data: List[Dict[str, Any]]) -> bytes:
        """Generate Excel comparison of all bids for client review"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Bid Comparison"
        
        ws['A1'] = f"Project: {project.name}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"Bid Comparison Report"
        ws['A2'].font = Font(bold=True, size=12)
        
        headers = ["Rank", "Contractor", "Company", "Total Amount", "Status", "Submitted Date"]
        header_row = 4
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = self.header_font
            cell.border = self.full_border
            cell.alignment = Alignment(horizontal='center')
        
        current_row = header_row + 1
        for bid in bids_data:
            ws.cell(row=current_row, column=1, value=bid.get('rank', 'N/A'))
            ws.cell(row=current_row, column=2, value=bid.get('contractor_name', ''))
            ws.cell(row=current_row, column=3, value=bid.get('contractor_company', ''))
            
            amount_cell = ws.cell(row=current_row, column=4, value=bid.get('total_amount', 0))
            amount_cell.number_format = '0.00'
            
            ws.cell(row=current_row, column=5, value=bid.get('status', ''))
            ws.cell(row=current_row, column=6, value=bid.get('submitted_at', ''))
            
            if bid.get('rank') == 1:
                for col in range(1, 7):
                    ws.cell(row=current_row, column=col).font = Font(bold=True)
            
            current_row += 1
        
        column_widths = [8, 20, 25, 15, 12, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
        
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        return excel_buffer.getvalue()
